"""Convert report/report.md (with Markdown tables and PNG figures) into report/report.docx."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

IMG_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
HEADING_RE = re.compile(r"^(?P<hashes>#{1,4})\s+(?P<title>.+)$")
INLINE_TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def _append_formatted_runs(paragraph, text: str) -> None:
    """Render **bold** and `code` inline spans cleanly into a python-docx paragraph."""
    # Strip markdown links [label](url) -> label for clean print typography
    cleaned = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    parts = INLINE_TOKEN_RE.split(cleaned)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def _render_table(doc: Document, table_lines: List[str]) -> None:
    rows: List[List[str]] = []
    for raw in table_lines:
        stripped = raw.strip().strip("|")
        cells = [c.strip() for c in stripped.split("|")]
        # Skip Markdown separator line like | :--- | :---: |
        if all(re.match(r"^:?-{2,}:?$", c) for c in cells if c != ""):
            continue
        rows.append(cells)

    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row_vals in enumerate(rows):
        row_cells = table.rows[r_idx].cells
        for c_idx in range(n_cols):
            val = row_vals[c_idx] if c_idx < len(row_vals) else ""
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _append_formatted_runs(p, val)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
            else:
                for run in p.runs:
                    if run.font.size is None:
                        run.font.size = Pt(9.0)


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    """Parse `md_path` and write formatted Word document to `docx_path`."""
    if not md_path.exists():
        raise FileNotFoundError(f"Markdown report not found: {md_path}")

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    base_dir = md_path.parent

    in_code = False
    code_buf: List[str] = []
    table_buf: List[str] = []

    def _flush_table() -> None:
        nonlocal table_buf
        if table_buf:
            _render_table(doc, table_buf)
            doc.add_paragraph()
            table_buf = []

    for line in lines:
        if line.strip().startswith("```"):
            _flush_table()
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9.0)
                code_buf = []
            continue

        if in_code:
            code_buf.append(line)
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_buf.append(line)
            continue
        else:
            _flush_table()

        stripped = line.strip()
        if not stripped or stripped == "---":
            continue

        m_head = HEADING_RE.match(stripped)
        if m_head:
            level = len(m_head.group("hashes"))
            title_text = re.sub(r"`([^`]+)`", r"\1", m_head.group("title").strip())
            h = doc.add_heading(title_text, level=min(level, 4))
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        m_img = IMG_RE.match(stripped)
        if m_img:
            alt = m_img.group("alt").strip()
            src_raw = m_img.group("src").strip()
            img_path = (base_dir / src_raw).resolve()
            if not img_path.exists():
                img_path = (base_dir.parent / src_raw).resolve()
            if img_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(str(img_path), width=Inches(6.3))
                if alt:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_cap = p_cap.add_run(alt)
                    r_cap.italic = True
                    r_cap.font.size = Pt(9.5)
                    r_cap.font.color.rgb = RGBColor(70, 70, 70)
            continue

        if stripped.startswith(">"):
            quote_txt = re.sub(r"^>\s*(\[!(?:NOTE|IMPORTANT|WARNING|TIP|CAUTION)\]\s*)?", "", stripped)
            if quote_txt:
                p_q = doc.add_paragraph()
                p_q.paragraph_format.left_indent = Inches(0.35)
                _append_formatted_runs(p_q, quote_txt)
                for r in p_q.runs:
                    r.italic = True
            continue

        m_ul = re.match(r"^[-*]\s+(.*)$", stripped)
        if m_ul:
            p_li = doc.add_paragraph(style="List Bullet")
            _append_formatted_runs(p_li, m_ul.group(1))
            continue

        m_ol = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m_ol:
            p_ol = doc.add_paragraph(style="List Number")
            _append_formatted_runs(p_ol, m_ol.group(1))
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        _append_formatted_runs(p, stripped)

    _flush_table()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Exported Word report -> {docx_path} ({docx_path.stat().st_size:,} bytes)")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build report.docx from report.md")
    parser.add_argument("--input", default="report/report.md")
    parser.add_argument("--output", default="report/report.docx")
    args = parser.parse_args()
    convert_markdown_to_docx(Path(args.input), Path(args.output))


if __name__ == "__main__":
    main()
