"""Unit test for Markdown-to-DOCX report converter."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from scripts.build_report_docx import convert_markdown_to_docx


def test_convert_markdown_to_docx_tables_and_headings(tmp_path: Path) -> None:
    md_file = tmp_path / "sample_report.md"
    docx_file = tmp_path / "sample_report.docx"
    md_file.write_text(
        "# Báo Cáo Thực Nghiệm\n\n"
        "## 1. Giới Thiệu\n"
        "Đoạn văn kiểm tra **in đậm** và `inline_code`.\n\n"
        "| Tham số | Giá trị |\n"
        "| :--- | :---: |\n"
        "| `ooo_ratio` | **50%** |\n"
        "| `batch_size` | 1000 |\n",
        encoding="utf-8",
    )
    convert_markdown_to_docx(md_file, docx_file)
    assert docx_file.exists() and docx_file.stat().st_size > 5_000
    doc = Document(str(docx_file))
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text == "ooo_ratio"
